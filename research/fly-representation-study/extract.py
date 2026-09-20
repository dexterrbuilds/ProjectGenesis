"""Read-only published XLSX extraction; bundled Python + openpyxl."""
import hashlib,json
from pathlib import Path
import openpyxl
H=Path(__file__).resolve().parent
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 assert (H/'FIT_FROZEN.json').exists(),'Freeze fit before reading challenge measurements'
 out=H/'data/challenges.json';assert not out.exists()
 w=openpyxl.load_workbook(H/'data/apl-fig3.xlsx',data_only=True,read_only=True);s=w['Sheet1'];rows=[]
 for label,col,n,stim in [('853',1,5,"alpha_prime"),('mb247',4,5,'alpha'),('c739',7,5,'alpha'),('np3061',10,6,'alpha')]:
  for row in range(3,3+n):
   a=float(s.cell(row,col).value);b=float(s.cell(row,col+1).value)
   rows.append({'driver':label,'source_row':row,'alpha':a,'alpha_prime':b,'stimulated':stim,
      'on':b if stim=='alpha_prime' else a,'off':a if stim=='alpha_prime' else b})
 w=openpyxl.load_workbook(H/'data/mbon14-egfp.xlsx',data_only=True,read_only=True);s=w['Tabelle1']
 mb=[{'cell':s.cell(i,1).value,'Vm_mV':float(s.cell(i,2).value),'tau_ms':float(s.cell(i,3).value),'C_pF':float(s.cell(i,4).value)} for i in range(4,8)]
 out.write_text(json.dumps({'apl':rows,'mbon14':mb,'mbon14_reported_mean':{'Vm_mV':s.cell(9,2).value,'tau_ms':s.cell(9,3).value,'C_pF':s.cell(9,4).value},
  'sources':[{'file':name,'sha256':sha(H/'data'/name),'url':url,'license':'CC-BY 4.0','attribution':who} for name,url,who in [
   ('apl-fig3.xlsx','https://cdn.elifesciences.org/articles/56954/elife-56954-fig3-data1-v2.xlsx','Amin et al. 2020, Figure 3I, eLife 56954'),
   ('mbon14-egfp.xlsx','https://cdn.elifesciences.org/articles/77578/elife-77578-table1-data1-v4.xlsx','Hafez et al. 2023, Table 1 source data 1, eLife 77578')]],
  'fit_before_extraction_sha256':sha(H/'FIT_FROZEN.json'),'negative_observations_retained':True,
  'quality_flags':['MBON14 supplementary displayed mean does not equal mean of four displayed cells.',
   'MBON14 supplementary cell 3 Vm/tau/C exactly matches original Table 1 cell 3; independence unresolved.',
   'APL Figure 3 is different indicator, stimulation and region pairing from Figure 4; parameter transfer is a challenge, not assumed valid.']},indent=2)+'\n')
if __name__=='__main__':main()
