"""Read-only extraction of CC-BY source tables. Run with bundled Python.

No model or Stage outcomes are inputs. Original workbook remains byte-for-byte intact.
"""
import hashlib
import json
from pathlib import Path
import openpyxl

H = Path(__file__).resolve().parent

def main():
    p = H/'data/amin-fig4.xlsx'
    w = openpyxl.load_workbook(p, data_only=True, read_only=True)
    s = w['Ark1']
    rows = []
    for i in range(4, 15):
        v = [s.cell(i, j).value for j in range(1, 5)]
        assert all(isinstance(x, (int, float)) for x in v)
        rows.append({'source_row': i, 'values': v})
    for split, parity in [('train', 0), ('validation', 1)]:
        target = H/'data'/f'apl-{split}.json'
        assert not target.exists(), 'Refuse to overwrite an extracted split'
        target.write_text(json.dumps([r for i,r in enumerate(rows) if i%2 == parity], indent=2)+'\n')
    controls = {label: [[s.cell(i,j).value for j in range(1,5)] for i in span]
                for label,span in [('KC_stimulation',range(18,29)),('negative_control',range(32,40))]}
    (H/'data/apl-controls.json').write_text(json.dumps(controls,indent=2)+'\n')
    meta = {'source': 'https://cdn.elifesciences.org/articles/56954/elife-56954-fig4-data1-v2.xlsx',
            'article': 'https://elifesciences.org/articles/56954', 'license': 'CC-BY 4.0',
            'source_sha256': hashlib.sha256(p.read_bytes()).hexdigest(),
            'sheet': 'Ark1', 'range': 'A4:D14', 'panel': '4D',
            'columns': ['calyx response to calyx stimulation','lobe response to calyx stimulation',
                        'calyx response to lobe stimulation','lobe response to lobe stimulation'],
            'measurement': 'mean deltaF/F over the figure-defined five-second window',
            'n_recordings': 11, 'n_flies_reported': 9,
            'fly_id_available': False, 'sex': 'not stated in retrieved methods',
            'split': 'source row indices 0,2,4,6,8,10 train; 1,3,5,7,9 validation',
            'exclusions': [], 'negative_observations_retained': True,
            'limitations': ['Split is recording-level, not demonstrably animal-independent.',
                           'On-site signal used as noisy predictor; not a voltage or conductance fit.',
                           'No cross-study validation; no FAFB individual was physiologically recorded.']}
    (H/'data/PROVENANCE.json').write_text(json.dumps(meta,indent=2)+'\n')

if __name__ == '__main__': main()
